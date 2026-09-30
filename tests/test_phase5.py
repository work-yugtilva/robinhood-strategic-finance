import csv
import hashlib
import json
import math
import re
import subprocess
import unittest
from pathlib import Path

import fitz
from openpyxl import load_workbook
from pptx import Presentation

from src.config import ROOT
from src.forecast_analysis import make_sensitivities, make_summaries
from src.forecast_model import SCENARIOS, actual_rows, build_assumptions, load_history, project
from src.phase4 import HISTORICAL_PATHS
from src.strategic_analysis import audit_interest, audit_margin, strategy_thresholds


class Phase5Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.f=json.loads((ROOT/'data/strategy/headline_facts.json').read_text())
        cls.h=load_history();_,cls.values=build_assumptions(cls.h)
        cls.actual=actual_rows(cls.h)
        cls.base=project(cls.h,cls.values,'Base')

    def test_interest_audit_reconciles_and_excludes_sweep_policy(self):
        audit=audit_interest(self.base,self.values)
        parts=audit['components_m']
        self.assertAlmostEqual(sum(parts.values()),audit['uniform_yield_100bp_m'],places=7)
        self.assertAlmostEqual(audit['uniform_yield_100bp_m'],898.8676988259105,places=5)
        self.assertAlmostEqual(parts['cash_sweep'],400.9015748031496,places=5)
        self.assertAlmostEqual(audit['policy_100bp_m'],264.69438769891525,places=5)
        self.assertEqual(audit['company_disclosed_2026_100bp_m'],350)
        for scenario in SCENARIOS:
            for period,inputs in self.values[scenario].items():
                if 'Q' in period:self.assertEqual(inputs['cash_sweep_rate_pass_through'],0)

    def test_margin_audit_and_thresholds(self):
        margin=audit_margin(self.actual,self.base)
        for row in margin:
            self.assertAlmostEqual(row['operating_expense_m'],sum(row[k] for k in
                ('brokerage_transaction_expense_m','technology_development_expense_m',
                 'operations_expense_m','credit_loss_expense_m','marketing_expense_m',
                 'general_administrative_expense_m')))
            self.assertAlmostEqual(row['operating_income_m'],row['revenue_m']-row['operating_expense_m'])
        thresholds=strategy_thresholds(self.h,self.values,self.base,margin)
        self.assertEqual(len(thresholds),5)
        target=margin[2]['operating_income_m']
        for row in thresholds[:3]:
            if row['required_input'] is not None:
                self.assertAlmostEqual(row['result_2027_oi_m'],target,places=6)
        self.assertIsNone(thresholds[1]['required_input'])

    def test_workbook_matches_model_and_has_no_formula_errors(self):
        path=ROOT/'models/Robinhood_Strategic_Finance_Model.xlsx'
        values=load_workbook(path,data_only=True,read_only=True)
        formulas=load_workbook(path,data_only=False,read_only=True)
        self.assertEqual(values.sheetnames,['Cover','Sources','Assumptions','Historical Financials',
            'Operating KPIs','Revenue Drivers','Expense Build','Forecast','Scenarios','Sensitivity',
            'Strategic Analysis','Executive Dashboard'])
        expected=self.f['model']['annual']['Base']['2027']
        self.assertAlmostEqual(values['Executive Dashboard']['E9'].value,expected['revenue_m'],places=6)
        self.assertAlmostEqual(values['Executive Dashboard']['E15'].value,expected['operating_margin_pct']/100,places=6)
        self.assertAlmostEqual(values['Forecast']['J10'].value,self.base[2]['revenue_m'],places=6)
        self.assertAlmostEqual(values['Strategic Analysis']['F22'].value,
                               self.f['model']['thresholds'][0]['required_input'],places=5)
        self.assertEqual(values['Assumptions']['D4'].value,1)
        self.assertTrue(str(formulas['Forecast']['J10'].value).startswith('='))
        self.assertTrue(str(formulas['Assumptions']['H8'].value).startswith('='))
        errors={'#REF!','#DIV/0!','#VALUE!','#NAME?','#NUM!','#NULL!','#N/A'}
        for ws in values:
            for row in ws:
                for cell in row:
                    self.assertNotIn(cell.value,errors,msg=f'{ws.title}!{cell.coordinate}') if isinstance(cell.value,str) else None
        for ws in formulas:
            for row in ws:
                for cell in row:
                    if cell.data_type=='f':self.assertNotIn('#REF!',cell.value)

    def test_deck_and_executive_text_use_same_figures(self):
        pptx=Presentation(ROOT/'presentation/Robinhood_Strategic_Finance_Executive_Deck.pptx')
        self.assertEqual(len(pptx.slides),8)
        deck='\n'.join(shape.text for slide in pptx.slides for shape in slide.shapes if shape.has_text_frame)
        with fitz.open(ROOT/'presentation/Robinhood_Strategic_Finance_Executive_Deck.pdf') as pdf:
            self.assertEqual(len(pdf),8)
            pdf_text='\n'.join(page.get_text() for page in pdf)
        b=self.f['model']['annual']['Base']['2027']
        rev=f'${b["revenue_m"]/1000:.2f}b'; margin=f'{b["operating_margin_pct"]:.1f}%'
        for text in (deck,pdf_text,(ROOT/'reports/MANAGEMENT_MEMO.md').read_text(),
                     (ROOT/'README.md').read_text(),(ROOT/'reports/INTERVIEW_GUIDE.md').read_text()):
            self.assertIn(rev,text)
            self.assertIn(margin,text)
        self.assertIn('not company guidance',deck)
        self.assertTrue(all((ROOT/'assets/screenshots'/p).exists() for p in
            ('executive_dashboard.png','historical_financials.png','revenue_driver_analysis.png',
             'forecast.png','sensitivity_analysis.png','strategic_analysis.png','executive_deck.png')))
        readme=(ROOT/'README.md').read_text()
        for image in re.findall(r'!\[[^]]*\]\((assets/screenshots/[^)]+)\)',readme):
            self.assertTrue((ROOT/image).exists(),msg=f'Missing README screenshot: {image}')

    def test_sources_status_and_language(self):
        f=self.f
        self.assertEqual(f['information_cutoff'],'2026-09-28')
        self.assertIn('sec.gov',f['sources']['q2_2026_filing'])
        addendum=(ROOT/'data/strategy/source_addendum.csv').read_text()
        self.assertIn('2026Q2_100bp_company_rate_sensitivity_m',addendum)
        self.assertIn('350',addendum)
        with (ROOT/'data/forecast/forecast_quarterly.csv').open() as handle:
            forecast=list(csv.DictReader(handle))
        self.assertEqual(next(r for r in forecast if r['period']=='2026Q3' and r['scenario']=='Base')['status'],'FORECAST')
        for path in (ROOT/'README.md',ROOT/'reports/MANAGEMENT_MEMO.md',
                     ROOT/'reports/STRATEGIC_DECISION.md'):
            text=path.read_text().lower()
            self.assertNotRegex(text,r'price target|buy hood shares|sell hood shares|discounted cash flow')
        self.assertLess(len((ROOT/'reports/MANAGEMENT_MEMO.md').read_text().split()),750)

    def test_repeat_build_is_byte_identical_and_history_untouched(self):
        paths=[*sorted((ROOT/'data/strategy').glob('*')),*sorted((ROOT/'assets/screenshots').glob('*.png')),
               ROOT/'models/Robinhood_Strategic_Finance_Model.xlsx',
               *sorted((ROOT/'presentation').glob('*')),
               ROOT/'reports/MANAGEMENT_MEMO.md',ROOT/'reports/INTERVIEW_GUIDE.md',ROOT/'README.md']
        def hashes(xs):return {str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in xs}
        before=hashes(paths);history_before=hashes(HISTORICAL_PATHS)
        subprocess.run(['python','-m','src.phase5'],cwd=ROOT,check=True,stdout=subprocess.DEVNULL)
        self.assertEqual(before,hashes(paths))
        self.assertEqual(history_before,hashes(HISTORICAL_PATHS))


if __name__=='__main__':
    unittest.main()

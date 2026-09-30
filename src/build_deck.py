"""Create an eight-slide executive deck from the audited Phase 5 facts."""
from __future__ import annotations

import json
import subprocess
import zipfile
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from .config import ROOT

OUT = ROOT / 'presentation'
PPTX = OUT / 'Robinhood_Strategic_Finance_Executive_Deck.pptx'
PDF = OUT / 'Robinhood_Strategic_Finance_Executive_Deck.pdf'
NAVY=RGBColor(23,43,69); TEAL=RGBColor(8,126,139); DARK=RGBColor(29,39,51)
MUTED=RGBColor(92,107,121); PALE=RGBColor(233,241,245); WHITE=RGBColor(255,255,255)
RED=RGBColor(169,58,58)


def textbox(slide, x, y, w, h, message, size=16, color=DARK, bold=False, align=None):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=tf.margin_bottom=Inches(.02)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0];p.text=message
    if align:p.alignment=align
    p.space_after=Pt(0);p.line_spacing=1.1
    for run in p.runs:
        run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=color
    return box


def rect(slide,x,y,w,h,fill,rounding=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounding else MSO_SHAPE.RECTANGLE,
                                 Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=fill
    shape.line.fill.background()
    return shape


def base_slide(prs,num,headline,source):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide,0,0,13.33,.13,TEAL)
    textbox(slide,.55,.35,12.25,.85,headline,25,NAVY,True)
    rect(slide,.55,1.32,12.25,.02,PALE)
    textbox(slide,.55,7.12,11.8,.2,source,8,MUTED)
    textbox(slide,12.4,7.09,.4,.23,str(num),9,MUTED,align=PP_ALIGN.RIGHT)
    return slide


def label_value(slide,x,y,label,value,sub='',accent=TEAL):
    rect(slide,x,y,3.7,1.25,PALE,True)
    textbox(slide,x+.17,y+.1,3.35,.25,label,11,MUTED)
    textbox(slide,x+.17,y+.39,3.35,.47,value,23,accent,True)
    if sub:textbox(slide,x+.17,y+.96,3.35,.2,sub,9,MUTED)


def bullet(slide,x,y,w,text,size=16,color=DARK):
    rect(slide,x,y+.11,.08,.08,TEAL)
    textbox(slide,x+.22,y-.02,w-.22,.65,text,size,color)


def bar(slide,x,y,w,label,value,max_value,color=TEAL,value_label=None):
    textbox(slide,x,y,2.0,.3,label,12,DARK)
    rect(slide,x+2.05,y+.03,w,.22,PALE)
    rect(slide,x+2.05,y+.03,w*max(0,value)/max_value,.22,color)
    textbox(slide,x+2.13+w,y-.01,1.6,.32,value_label or f'{value:,.0f}',12,DARK,True)


def canonical_zip(target):
    tmp=target.with_suffix('.stable')
    with zipfile.ZipFile(target) as src,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as dst:
        for name in sorted(src.namelist()):
            info=zipfile.ZipInfo(name,date_time=(1980,1,1,0,0,0))
            info.compress_type=zipfile.ZIP_DEFLATED
            dst.writestr(info,src.read(name))
    tmp.replace(target)


def run():
    f=json.loads((ROOT/'data/strategy/headline_facts.json').read_text())
    m=f['model'];a=m['annual'];r=f['reported'];sh=m['shocks'];t=m['thresholds']
    b27=a['Base']['2027'];b28=a['Base']['2028'];u27=a['Upside']['2027'];d27=a['Downside']['2027']
    spec=[
      {'headline':'Revenue quality depends on trading engagement, earning balances, and expense conversion',
       'question':'What is the management conclusion?'},
      {'headline':'Revenue rose 52% in 2025 as customers and assets expanded; expense commitments are rising',
       'question':'What changed financially?'},
      {'headline':'Measured trading activity drove most 2024 transaction growth, with later scope less comparable',
       'question':'What actually drove growth?'},
      {'headline':'Assets grew faster than customers, but acquisitions and market moves limit organic inference',
       'question':'Is scale translating into monetization?'},
      {'headline':f'Base reaches ${b27["revenue_m"]/1000:.2f}b 2027 revenue at {b27["operating_margin_pct"]:.1f}% margin',
       'question':'What must happen operationally?'},
      {'headline':'Transaction pricing and rate pass-through can move 2027 profit by hundreds of millions',
       'question':'What can break the forecast?'},
      {'headline':'Each growth path needs a different operating condition; none has public-data initiative ROI',
       'question':'Which paths deserve attention?'},
      {'headline':'Protect monetization, deepen earning balances, and gate spend to realized profit conversion',
       'question':'What should management monitor?'},
    ]
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'slide_spec.json').write_text(json.dumps(spec,sort_keys=True,indent=2)+'\n')
    prs=Presentation();prs.slide_width=Inches(13.33);prs.slide_height=Inches(7.5)
    prs.core_properties.author='Independent portfolio project'
    prs.core_properties.title='Robinhood Strategic Finance & Operating Model'
    prs.core_properties.subject='Public-data strategic finance analysis; cutoff 2026-09-28'
    # 1: conclusion
    s=base_slide(prs,1,spec[0]['headline'],'Sources: SEC filings; Robinhood releases; audited analyst model. Forecasts are not company guidance.')
    label_value(s,.65,1.72,'2027 BASE REVENUE',f'${b27["revenue_m"]/1000:.2f}b','Analyst forecast')
    label_value(s,4.83,1.72,'2027 OPERATING MARGIN',f'{b27["operating_margin_pct"]:.1f}%','GAAP analyst forecast')
    label_value(s,9.0,1.72,'2027 SCENARIO RANGE',f'${d27["revenue_m"]/1000:.2f}–{u27["revenue_m"]/1000:.2f}b','Downside to Upside')
    bullet(s,.8,3.45,11.8,f'Measured 2024 comparable-product activity contributed ${f["derived"]["2024_activity_bridge_m"]:,.0f}m of ${f["derived"]["2024_transaction_growth_m"]:,.0f}m transaction growth.',17)
    bullet(s,.8,4.35,11.8,'2027 Base interest/balance revenue is 36% of total; disclosed Gold is only 4.5%.',17)
    bullet(s,.8,5.25,11.8,'Base margin depends on revenue rising faster than the six-category expense plan.',17)
    textbox(s,.82,6.35,11.9,.42,'Major risk: trading or crypto normalization alongside weaker realized yields.',13,RED,True)
    # 2: history
    s=base_slide(prs,2,spec[1]['headline'],'Sources: FY 2025 10-K/A; Q2 2026 10-Q; Phase 3 historical dataset.')
    for y,v in [('2023',r['2023_revenue_m']),('2024',r['2024_revenue_m']),('2025',r['2025_revenue_m'])]:
        bar(s,.75,1.85+(int(y)-2023)*.82,5.7,y,v,5000,value_label=f'${v/1000:.2f}b')
    label_value(s,.75,4.75,'2026 Q2 REVENUE',f'${r["2026Q2_revenue_m"]/1000:.2f}b','Reported quarter')
    label_value(s,4.85,4.75,'FUNDED CUSTOMERS',f'{r["2026Q2_funded_customers_m"]:.1f}m','2026 Q2 reported')
    label_value(s,8.95,4.75,'PLATFORM ASSETS',f'${r["2026Q2_platform_assets_b"]:.1f}b','2026 Q2 reported')
    textbox(s,.78,6.25,11.7,.43,'2026 is two reported quarters plus two forecast quarters; it is not a reported full year.',12,MUTED)
    # 3: historical driver attribution
    s=base_slide(prs,3,spec[2]['headline'],'Sources: Phase 3 exact transaction-revenue bridge; 2025 10-K/A.')
    activity=f['derived']['2024_activity_bridge_m']; total=f['derived']['2024_transaction_growth_m']
    bar(s,.75,1.85,5.8,'Activity',activity,total,value_label=f'${activity:,.0f}m')
    bar(s,.75,2.75,5.8,'Other effects',total-activity,total,color=MUTED,value_label=f'${total-activity:,.0f}m')
    textbox(s,.8,3.83,11.7,.46,f'${total:,.0f}m total 2024 vs 2023 transaction-revenue increase; bridge reconciles exactly.',16,NAVY,True)
    bullet(s,.8,4.78,11.9,'Customer scale and monetization contribute, but measured activity dominates this comparison.',15)
    bullet(s,.8,5.57,11.9,'The 2025 bridge retains new products and acquisitions in an explicit residual.',15)
    # 4: scale monetization
    s=base_slide(prs,4,spec[3]['headline'],'Sources: FY 2025 10-K/A; Phase 3 monetization table; acquisition adjustments described in structural breaks.')
    label_value(s,.72,1.78,'CUSTOMERS',f'{r["2023Q4_funded_customers_m"]:.1f}m → {r["2025Q4_funded_customers_m"]:.1f}m','2023 Q4 to 2025 Q4')
    label_value(s,4.81,1.78,'PLATFORM ASSETS',f'${r["2023Q4_platform_assets_b"]:.0f}b → ${r["2025Q4_platform_assets_b"]:.0f}b','Revised 2025 Q4: $322.1b')
    label_value(s,8.9,1.78,'ASSETS / CUSTOMER',f'${1000*r["2023Q4_platform_assets_b"]/r["2023Q4_funded_customers_m"]:,.0f} → ${1000*r["2025Q4_platform_assets_b"]/r["2025Q4_funded_customers_m"]:,.0f}','Derived; rounded')
    bullet(s,.8,3.58,11.9,'Customer assets rose more rapidly than funded customers, creating more balance potential.',16)
    bullet(s,.8,4.51,11.9,'Reported asset growth includes acquisitions and market value changes; it is not organic deposits.',16)
    bullet(s,.8,5.45,11.9,'The forecast links only captured earning balances and activity to revenue, not all platform assets.',16)
    # 5: forecast
    s=base_slide(prs,5,spec[4]['headline'],'Sources: audited Base/Upside/Downside operating model; 2026 is mixed actual/forecast.')
    for j,(name,series,color) in enumerate([('Downside',a['Downside'],MUTED),('Base',a['Base'],TEAL),('Upside',a['Upside'],NAVY)]):
        bar(s,.8,1.86+j*.78,5.6,name,series['2027']['revenue_m'],8000,color,value_label=f'${series["2027"]["revenue_m"]/1000:.2f}b')
        textbox(s,10.1,1.84+j*.78,2,.28,f'{series["2027"]["operating_margin_pct"]:.1f}% margin',13,color,True)
    textbox(s,.81,4.55,11.65,.48,f'Base 2028: ${b28["revenue_m"]/1000:.2f}b revenue at {b28["operating_margin_pct"]:.1f}% operating margin.',16,NAVY,True)
    bullet(s,.8,5.48,11.8,'Base keeps normal-quarter additions near 0.3m and deposits near $17.7b.',15)
    bullet(s,.8,6.12,11.8,'Options activity and balances carry growth; expense growth remains below revenue growth.',15)
    # 6: sensitivity and mix
    s=base_slide(prs,6,spec[5]['headline'],'Sources: defined 2027 Base shocks; rate audit against June 2026 10-Q disclosure.')
    label_value(s,.75,1.72,'TRANSACTION PRICE -20%',f'-${abs(sh["transaction_monetization_down_20pct_m"]):,.0f}m','2027 revenue effect',RED)
    label_value(s,4.83,1.72,'POLICY RATE -100 BP',f'-${abs(sh["policy_rate_down_100bp_m"]):,.0f}m','2027 revenue effect',RED)
    label_value(s,8.91,1.72,'OPTIONS ACTIVITY -20%',f'-${abs(sh["options_activity_down_20pct_oi_m"]):,.0f}m','2027 operating-income effect',RED)
    mix=[('Market activity',b27['market_activity_pct'],TEAL),('Interest / balance',b27['interest_balance_pct'],NAVY),
         ('Gold',b27['disclosed_recurring_pct'],MUTED),('Unclassified',b27['unclassified_pct'],RGBColor(185,194,201))]
    x=.82
    for label,value,color in mix:
        rect(s,x,4.05,11.6*value/100,.35,color)
        x+=11.6*value/100
    textbox(s,.82,4.54,11.7,.43,'2027 Base mix: 52% market activity · 36% interest/balance · 4.5% Gold · 6.9% unclassified.',14,NAVY)
    textbox(s,.82,5.65,11.7,.75,'The $899m all-component +100 bp yield stress includes Cash Sweep. It is not the policy-rate case; Robinhood excludes sweep from its disclosed short-term rate sensitivity.',13,MUTED)
    # 7: paths and thresholds
    s=base_slide(prs,7,spec[6]['headline'],'Sources: audited decision thresholds and defined sensitivities; no product ROI is inferred.')
    for i,(name,impact,condition) in enumerate([
        ('Customer & asset scale',f'+${m["paths"][0]["2027_revenue_delta_m"]:,.0f}m revenue at +25% net additions','Additional customers must monetize beyond near-term marketing costs.'),
        ('Transaction engagement',f'+${m["paths"][1]["2027_revenue_delta_m"]:,.0f}m revenue at +20% options units',f'{t[0]["required_input"]:.1f} options contracts/customer offsets a modeled -50 bp rate case.'),
        ('Balance / Gold monetization',f'+${m["paths"][2]["2027_revenue_delta_m"]:,.0f}m revenue at +50% cash capture','Net yield and captured earning balances must persist; Gold remains modest.')]):
        y=1.75+i*1.55;rect(s,.72,y,11.85,1.35,PALE,True)
        textbox(s,.94,y+.12,4.0,.38,name,16,NAVY,True)
        textbox(s,5.08,y+.12,7.16,.35,impact,15,TEAL,True)
        textbox(s,.94,y+.70,11.2,.39,condition,13,DARK)
    textbox(s,.8,6.55,11.8,.28,f'A ${t[2]["required_input"]:.1f}b quarterly deposit threshold to offset -10% transaction pricing is outside recent history.',12,RED,True)
    # 8: management implications
    s=base_slide(prs,8,spec[7]['headline'],'Sources: management memo, CFO challenge, audited model. Independent portfolio analysis.')
    for i,p in enumerate(f['recommendation']['priorities']):bullet(s,.78,1.7+i*.76,11.8,f'{i+1}. {p}',15)
    textbox(s,.8,4.28,3.0,.3,'Two risks',14,NAVY,True)
    for i,p in enumerate(f['recommendation']['risks']):bullet(s,.8,4.69+i*.65,11.7,p,13,RED)
    textbox(s,.8,6.16,11.7,.58,'Monitor additions; deposit capture; options units/pricing; crypto bps; net yield and expense/revenue.',13,TEAL,True)
    prs.save(PPTX);canonical_zip(PPTX)
    subprocess.run(['soffice','-env:UserInstallation=file:///tmp/robinhood-phase5-office',
                    '--headless','--convert-to','pdf','--outdir',str(OUT),str(PPTX)],
                   check=True,capture_output=True,text=True)
    if not PDF.exists():raise RuntimeError('Deck PDF conversion did not create output')
    document=fitz.open(PDF)
    if len(document)!=8:raise RuntimeError('Deck PDF does not have eight slides')
    metadata={'title':'Robinhood Strategic Finance & Operating Model',
              'author':'Independent portfolio project','subject':'Public-data strategic finance analysis',
              'creationDate':'D:20260928000000','modDate':'D:20260928000000'}
    document.set_metadata(metadata)
    document.del_xml_metadata()
    document.xref_set_key(-1,'ID','[<526F62696E686F6F6446696E616E6365> <526F62696E686F6F6446696E616E6365>]')
    stable=PDF.with_suffix('.stable.pdf')
    document.save(stable,garbage=4,deflate=True,no_new_id=1)
    document.close();stable.replace(PDF)
    document=fitz.open(PDF)
    screenshot=ROOT/'assets/screenshots/executive_deck.png'
    screenshot.parent.mkdir(parents=True,exist_ok=True)
    document[0].get_pixmap(matrix=fitz.Matrix(1.25,1.25),alpha=False).save(screenshot)
    document.close()
    print('Eight-slide PPTX and PDF exported.')


if __name__=='__main__':
    run()
